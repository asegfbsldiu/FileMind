"""Opt-in organization content, mock API payloads, and safe review integration."""

import base64
import io
import json
from pathlib import Path
from threading import Event
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from PIL import Image
from PySide6.QtTest import QTest
from PySide6.QtCore import Qt
from PySide6.QtGui import QTextCursor
from PySide6.QtWidgets import QMessageBox

from filemind.ai.content import MAX_EXCERPT_CHARS, prepare_content_metadata
from filemind.ai.openai_compatible import OpenAICompatibleClient, _make_batches
from filemind.ai.prompts import build_organize_content
from filemind.ai.schemas import AIOrganizationResult, file_metadata
from filemind.core.scanner import scan_files
from filemind.models import FileSuggestion
from filemind.storage.settings import AISettings
from filemind.ui.main_window import MainWindow
from filemind.ui.pages.organize_page import OrganizePage
from tools.review_ui import build_review
from tests.test_extended_conversion import create_text_pdf


SETTINGS = AISettings("deepseek", "https://api.deepseek.com", "deepseek-flash", 10, "fixture-only")


def prepare(root: Path, settings: AISettings = SETTINGS):
    """Use generated fixtures only; no user configuration or network."""
    return prepare_content_metadata(root, scan_files(root), settings, Event())


def test_metadata_default_does_not_read_contents(tmp_path: Path) -> None:
    source = tmp_path / "generated.txt"
    source.write_text("private body", encoding="utf-8")
    with patch.object(Path, "read_bytes", side_effect=AssertionError("no content read")):
        item = file_metadata(scan_files(tmp_path)[0])
        payload = build_organize_content([item], "归类", True)
    assert "private body" not in payload
    assert item.content_excerpt == "" and item.image_data_url == ""


def test_content_text_is_bounded_and_obvious_credential_values_redacted(tmp_path: Path) -> None:
    (tmp_path / "abc123.txt").write_text("数据结构 二叉树\napi_key=sk-fixture123456789\n" + "课程" * 3000, encoding="utf-8")
    files, failures = prepare(tmp_path)
    assert not failures
    assert "数据结构" in files[0].content_excerpt
    assert len(files[0].content_excerpt) <= MAX_EXCERPT_CHARS
    assert "sk-fixture" not in files[0].content_excerpt
    assert "[已遮蔽凭据]" in files[0].content_excerpt
    assert str(tmp_path) not in json.dumps(build_organize_content(files, "归类", True), ensure_ascii=False)


@pytest.mark.parametrize("name", [".env", "api_key.txt", "密码.txt", "id_rsa", "private.pem", "秘钥.txt", "sk_fixture.txt"])
def test_credential_names_are_excluded_before_any_file_is_opened(tmp_path: Path, name: str) -> None:
    (tmp_path / name).write_text("fixture secret", encoding="utf-8")
    with patch.object(Path, "open", side_effect=AssertionError("protected file read")):
        files, failures = prepare(tmp_path)
    assert files == []
    assert len(failures) == 1 and "未读取" in failures[0].message


def test_image_payload_has_stable_id_and_preserves_aspect_without_exif(tmp_path: Path) -> None:
    with Image.new("RGB", (2000, 1000), "navy") as image:
        image.save(tmp_path / "abc123.png")
    files, failures = prepare(tmp_path)
    assert not failures
    content = build_organize_content(files, "识别图片并命名", True)
    assert files[0].file_id in content[1]["text"]
    data_url = content[2]["image_url"]["url"]
    with Image.open(io.BytesIO(base64.b64decode(data_url.split(",", 1)[1]))) as image:
        assert image.size == (1024, 512)
        assert not image.getexif()


@pytest.mark.parametrize("kind", ["pdf", "docx", "pptx", "xlsx"])
def test_generated_document_excerpt_reaches_metadata(tmp_path: Path, kind: str) -> None:
    source = tmp_path / f"opaque.{kind}"
    expected = "Generated PDF page 1" if kind == "pdf" else "Binary tree course"
    if kind == "pdf":
        create_text_pdf(source, page_count=4)
    elif kind == "docx":
        from docx import Document
        document = Document()
        document.add_paragraph(expected)
        document.save(source)
    elif kind == "pptx":
        from pptx import Presentation
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = expected
        document.save(source)
    else:
        from openpyxl import Workbook
        document = Workbook()
        document.active["A1"] = expected
        document.save(source)
        document.close()
    files, failures = prepare(tmp_path)
    assert not failures
    assert expected in files[0].content_excerpt
    assert "Generated PDF page 4" not in files[0].content_excerpt


def test_read_limit_marks_every_unprocessed_file_explicitly(tmp_path: Path) -> None:
    for index in range(3):
        (tmp_path / f"{index}.txt").write_text("generated text", encoding="utf-8")
    with patch("filemind.ai.content.MAX_CONTENT_FILES", 2):
        files, failures = prepare(tmp_path)
    assert len(files) == 2 and len(failures) == 1
    assert len({item.file_id for item in files} | set(failures[0].file_ids)) == 3


def test_unverified_image_model_is_explicit_failure(tmp_path: Path) -> None:
    with Image.new("RGB", (10, 10)) as image:
        image.save(tmp_path / "sample.png")
    settings = AISettings("custom", "https://example.invalid/v1", "text-only", 10, "fixture")
    files, failures = prepare(tmp_path, settings)
    assert not files
    assert "图片理解" in failures[0].message


def test_changed_source_and_cancelled_analysis_do_not_send(tmp_path: Path) -> None:
    source = tmp_path / "sample.txt"
    source.write_text("old", encoding="utf-8")
    scanned = scan_files(tmp_path)
    source.write_text("changed length", encoding="utf-8")
    files, failures = prepare_content_metadata(tmp_path, scanned, SETTINGS, Event())
    assert not files and "变化" in failures[0].message
    cancel = Event()
    cancel.set()
    with pytest.raises(InterruptedError):
        prepare_content_metadata(tmp_path, scanned, SETTINGS, cancel)
    assert source.read_text(encoding="utf-8") == "changed length"


def test_image_and_text_request_reaches_sdk_and_validated_name(tmp_path: Path) -> None:
    with Image.new("RGB", (20, 10), "white") as image:
        image.save(tmp_path / "f123.png")
    files, _ = prepare(tmp_path)
    row = files[0]
    response = {"suggestions": [{"file_id": row.file_id, "category": "图片", "target_subdirectory": "课程",
                                "new_name": "课程示意图.png", "reason": "图中内容"}]}
    with patch("filemind.ai.openai_compatible.OpenAI") as sdk:
        sdk.return_value.chat.completions.create.return_value = SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(response, ensure_ascii=False)))])
        result = OpenAICompatibleClient(SETTINGS).organize_files(files, "按内容命名", True)
        sent = sdk.return_value.chat.completions.create.call_args.kwargs["messages"][1]["content"]
    assert isinstance(sent, list) and sent[-1]["type"] == "image_url"
    assert result.suggestions[row.file_id].new_name == "课程示意图.png"
    assert not result.failures


def test_images_are_batched_by_count_and_text_by_size(tmp_path: Path) -> None:
    for index in range(9):
        with Image.new("RGB", (10, 10)) as image:
            image.save(tmp_path / f"{index}.png")
    files, failures = prepare(tmp_path)
    batches = _make_batches(files, "归类", True)
    assert not failures and [len(batch) for batch in batches] == [4, 4, 1]
    assert {item.file_id for batch in batches for item in batch} == {item.file_id for item in files}


def test_content_suggestion_enters_plan_while_excluded_file_cannot_execute(tmp_path: Path) -> None:
    (tmp_path / "abc123.txt").write_text("二叉树课程第1章", encoding="utf-8")
    (tmp_path / "密钥.txt").write_text("fixture only", encoding="utf-8")
    captured = []

    class Model:
        def __init__(self, _settings):
            pass

        def organize_files(self, files, *_args):
            captured.extend(files)
            row = files[0]
            return AIOrganizationResult({row.file_id: FileSuggestion(
                row.file_name, "课程", "二叉树第一章.txt", "开头片段标明课程", row.file_id)}, ())

    with patch("filemind.ui.main_window.OpenAICompatibleClient", Model):
        plan, _, _, failed = MainWindow._scan_and_plan(
            tmp_path, True, SETTINGS, "内容命名", True, lambda *_: None, Event(), content_analysis=True)
    assert len(captured) == 1 and "二叉树" in captured[0].content_excerpt
    renamed = next(item for item in plan.operations if item.original_name == "abc123.txt")
    protected = next(item for item in plan.operations if item.original_name == "密钥.txt")
    assert renamed.new_name == "二叉树第一章.txt" and renamed.classification.source == "AI"
    assert protected.warning and failed
    assert not renamed.target.exists()


def test_content_preset_enables_ai_and_rename_and_keeps_input_editable() -> None:
    page = OrganizePage()
    assert not page.content_checkbox.isChecked()
    page.preset_combo.setCurrentIndex(page.preset_combo.findData("content"))
    assert page.mode_combo.currentData() == "ai"
    assert page.content_checkbox.isChecked() and page.rename_checkbox.isChecked()
    assert "片段" in page.ai_privacy_label.text()
    page.rule_input.setPlainText("按内容识别")
    page.mode_combo.setCurrentIndex(0)
    page.rule_input.moveCursor(QTextCursor.MoveOperation.End)
    QTest.keyClicks(page.rule_input, " editable")
    assert page.rule_input.toPlainText().endswith(" editable")
    assert "不会被" in page.rule_mode_note.text()
    page.close()


def test_declining_content_consent_never_reads_contents_or_starts_network(tmp_path: Path) -> None:
    window, _ = build_review(tmp_path, "orbit")
    selected = tmp_path / "selected"
    selected.mkdir()
    (selected / "data.txt").write_text("generated", encoding="utf-8")
    page = window.organize_page
    page.set_directory(str(selected))
    page.preset_combo.setCurrentIndex(page.preset_combo.findData("content"))
    with patch.object(window, "_prepare_ai_settings", return_value=SETTINGS), \
            patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No), \
            patch("filemind.ai.content._read_snapshot", side_effect=AssertionError("not authorized")), \
            patch("filemind.ui.main_window.OpenAICompatibleClient", side_effect=AssertionError("no network")):
        page.rule_send_button.click()
    assert window._active_worker is None and page.plan is None
    assert "已取消内容识别" in page.page_status.text()
    window.close()
