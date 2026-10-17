"""Exercise extended local converters using generated, disposable fixtures."""

from __future__ import annotations

import csv
import io
import json
import subprocess
import wave
from pathlib import Path
from threading import Event
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from PIL import Image

from filemind.conversion import media
from filemind.conversion.base import ConversionError
from filemind.conversion.models import ConversionResult
from filemind.conversion.service import convert_file, inspect_input


def convert_generated(source: Path, target: str, output: Path) -> ConversionResult:
    """Run the public pipeline without UI, network or personal files."""
    return convert_file(inspect_input(source), target, output, lambda *_: None, Event())


def create_text_pdf(output: Path, page_count: int = 2) -> None:
    """Build a minimal real PDF with a standard font, independent of OS fonts."""
    page_ids = [4 + index * 2 for index in range(page_count)]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_index in range(page_count):
        stream = f"BT /F1 16 Tf 30 750 Td (Generated PDF page {page_index + 1}) Tj ET\n".encode()
        objects.append((
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_ids[page_index] + 1} 0 R >>"
        ).encode())
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"endstream")
    contents = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(contents))
        contents.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    xref = len(contents)
    contents.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        contents.extend(f"{offset:010} 00000 n \n".encode())
    contents.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    output.write_bytes(contents)


def assert_no_staging(folder: Path) -> None:
    """Every failed conversion must clean its own intermediate files."""
    assert not list(folder.glob(".filemind-convert-*"))


def test_pdf_text_and_page_images_use_real_document_content(tmp_path: Path) -> None:
    source = tmp_path / "generated.pdf"
    create_text_pdf(source)
    assert {"pdf_text", "pdf_png", "pdf_jpeg"} <= {
        target.key for target in inspect_input(source).targets
    }
    text_output = tmp_path / "pages.txt"
    convert_generated(source, "pdf_text", text_output)
    text = text_output.read_text(encoding="utf-8")
    assert "Generated PDF page 1" in text
    assert "Generated PDF page 2" in text
    image_output = tmp_path / "pages.zip"
    convert_generated(source, "pdf_png", image_output)
    with ZipFile(image_output) as archive:
        assert archive.namelist() == ["page-001.png", "page-002.png"]
        for entry in archive.namelist():
            with Image.open(io.BytesIO(archive.read(entry))) as image:
                assert image.format == "PNG"
                assert image.width > 500 and image.height > 500


def test_pdf_cancel_between_pages_does_not_publish_partial_archive(tmp_path: Path) -> None:
    source = tmp_path / "generated.pdf"
    create_text_pdf(source)
    cancel = Event()

    def stop_after_first_page(completed: int, total: int) -> None:
        if completed == 1 and total == 2:
            cancel.set()

    output = tmp_path / "cancelled.zip"
    with pytest.raises(InterruptedError):
        convert_file(inspect_input(source), "pdf_png", output, stop_after_first_page, cancel)
    assert not output.exists()
    assert_no_staging(tmp_path)


def test_scanned_pdf_text_extraction_reports_missing_text_layer(tmp_path: Path) -> None:
    source = tmp_path / "image.png"
    with Image.new("RGB", (30, 20), "white") as image:
        image.save(source)
    pdf = tmp_path / "scan.pdf"
    convert_generated(source, "pdf", pdf)
    output = tmp_path / "scan.txt"
    with pytest.raises(ConversionError, match="文本层"):
        convert_generated(pdf, "pdf_text", output)
    assert not output.exists()
    assert_no_staging(tmp_path)


def test_every_advertised_raster_encoder_produces_the_named_format(tmp_path: Path) -> None:
    source = tmp_path / "pixels.png"
    with Image.new("RGB", (32, 24), "navy") as image:
        image.save(source)
    expected = {"png": "PNG", "jpeg": "JPEG", "webp": "WEBP", "bmp": "BMP",
                "tiff": "TIFF", "gif": "GIF", "avif": "AVIF"}
    targets = [target for target in inspect_input(source).targets if target.key in expected]
    assert {"png", "jpeg", "gif", "tiff"} <= {target.key for target in targets}
    for target in targets:
        output = tmp_path / f"encoded{target.extension}"
        convert_generated(source, target.key, output)
        with Image.open(output) as image:
            image.load()
            assert image.format == expected[target.key]
            assert image.size == (32, 24)


def test_invalid_pdf_is_a_readable_conversion_error(tmp_path: Path) -> None:
    source = tmp_path / "broken.pdf"
    source.write_bytes(b"%PDF-1.7\nnot a document")
    with pytest.raises(ConversionError, match="PDF"):
        inspect_input(source)


def test_docx_text_and_html_preserve_paragraphs_and_tables(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "course.docx"
    document = Document()
    document.add_paragraph("课程：数据结构 <重点>")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "章节"
    table.cell(0, 1).text = "二叉树"
    document.save(source)
    output = tmp_path / "course.txt"
    convert_generated(source, "office_text", output)
    text = output.read_text(encoding="utf-8")
    assert "课程：数据结构 <重点>" in text
    assert "章节\t二叉树" in text
    html_output = tmp_path / "course.html"
    convert_generated(source, "office_html", html_output)
    assert "&lt;重点&gt;" in html_output.read_text(encoding="utf-8")


def test_pptx_extracts_text_in_slide_order(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "lecture.pptx"
    presentation = Presentation()
    for content in ("First generated slide", "第二页：链表"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text = content
    presentation.save(source)
    output = tmp_path / "lecture.txt"
    convert_generated(source, "office_text", output)
    text = output.read_text(encoding="utf-8")
    assert text.index("First generated slide") < text.index("第二页：链表")
    assert "第 1 页" in text and "第 2 页" in text


def test_xlsx_exports_all_sheets_without_executing_formulas(tmp_path: Path) -> None:
    from openpyxl import Workbook

    source = tmp_path / "all-sheets.xlsx"
    workbook = Workbook()
    workbook.active.title = "课程"
    workbook.active.append(["课程", "成绩"])
    workbook.active.append(["数据结构", 95])
    second = workbook.create_sheet("计算")
    second.append(["未计算的公式", "=1+1"])
    workbook.save(source)
    workbook.close()
    json_output = tmp_path / "sheets.json"
    convert_generated(source, "sheet_json", json_output)
    contents = json.loads(json_output.read_text(encoding="utf-8"))
    assert set(contents) == {"课程", "计算"}
    assert contents["课程"][1] == ["数据结构", 95]
    assert contents["计算"][0] == ["未计算的公式", None]
    csv_output = tmp_path / "sheets.zip"
    convert_generated(source, "sheet_csv", csv_output)
    with ZipFile(csv_output) as archive:
        assert len(archive.namelist()) == 2
        rows = [archive.read(name).decode("utf-8-sig") for name in archive.namelist()]
        assert any("数据结构,95" in text for text in rows)
        assert any("未计算的公式" in text for text in rows)


def test_csv_xlsx_keeps_formula_like_cells_as_literal_text(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    source = tmp_path / "literal.csv"
    values = ["=1+1", "+SUM(1)", "-1+2", "@SUM(2)", "00123"]
    with source.open("w", encoding="utf-8", newline="") as stream:
        csv.writer(stream).writerow(values)
    output = tmp_path / "literal.xlsx"
    convert_generated(source, "csv_xlsx", output)
    workbook = load_workbook(output, data_only=False)
    try:
        assert [cell.value for cell in workbook.active[1]] == values
        assert all(cell.data_type == "s" for cell in workbook.active[1])
    finally:
        workbook.close()


def test_text_html_docx_and_html_extraction_are_real(tmp_path: Path) -> None:
    from docx import Document

    source = tmp_path / "notes.txt"
    source.write_text("标题 <code>\n第二行", encoding="utf-8")
    html_output = tmp_path / "notes.html"
    convert_generated(source, "text_html", html_output)
    assert "&lt;code&gt;" in html_output.read_text(encoding="utf-8")
    docx_output = tmp_path / "notes.docx"
    convert_generated(source, "text_docx", docx_output)
    assert [paragraph.text for paragraph in Document(docx_output).paragraphs] == [
        "标题 <code>", "第二行"
    ]
    html_input = tmp_path / "visible.html"
    html_input.write_text(
        '<h1>标题</h1><script>secret_script()</script><style>.hidden{}</style>'
        '<p>可见 &amp; 文字</p><img src="https://invalid.example/image">', encoding="utf-8"
    )
    plain = tmp_path / "visible.txt"
    convert_generated(html_input, "html_text", plain)
    assert plain.read_text(encoding="utf-8") == "标题\n可见 & 文字"


def test_json_yaml_round_trip_and_xml_formatting(tmp_path: Path) -> None:
    source = tmp_path / "data.json"
    contents = {"课程": ["数据结构", "Java"], "启用": True, "数量": 2}
    source.write_text(json.dumps(contents, ensure_ascii=False), encoding="utf-8")
    yaml_output = tmp_path / "data.yaml"
    convert_generated(source, "json_yaml", yaml_output)
    json_output = tmp_path / "roundtrip.json"
    convert_generated(yaml_output, "yaml_json", json_output)
    assert json.loads(json_output.read_text(encoding="utf-8")) == contents
    xml_source = tmp_path / "data.xml"
    xml_source.write_text('<root><course name="数据结构">树</course></root>', encoding="utf-8")
    xml_output = tmp_path / "formatted.xml"
    convert_generated(xml_source, "xml_pretty", xml_output)
    assert "\n  <course" in xml_output.read_text(encoding="utf-8")
    assert "数据结构" in xml_output.read_text(encoding="utf-8")


@pytest.mark.parametrize("suffix,contents,target", [
    ("yaml", "value: !!python/object/apply:os.system ['do-not-run']", "yaml_json"),
    ("yaml", "value: &loop [*loop]", "yaml_json"),
    ("yaml", "date: 2026-10-16", "yaml_json"),
    ("xml", '<!DOCTYPE root [<!ENTITY value SYSTEM "file:///not-read">]><root>&value;</root>', "xml_pretty"),
    ("xml", "<root><unclosed></root>", "xml_pretty"),
])
def test_unsafe_or_invalid_structured_text_reports_conversion_error(
    tmp_path: Path, suffix: str, contents: str, target: str
) -> None:
    source = tmp_path / f"input.{suffix}"
    source.write_text(contents, encoding="utf-8")
    output = tmp_path / ("output.json" if suffix == "yaml" else "output.xml")
    with pytest.raises(ConversionError):
        convert_generated(source, target, output)
    assert not output.exists()
    assert_no_staging(tmp_path)


def create_epub(source: Path, chapter_href: str = "second.xhtml") -> None:
    """Build a minimal two-chapter EPUB with explicit spine order."""
    container = '<container><rootfiles><rootfile full-path="OPS/content.opf"/></rootfiles></container>'
    package = (
        '<package><manifest><item id="one" href="first.xhtml"/>'
        f'<item id="two" href="{chapter_href}"/></manifest>'
        '<spine><itemref idref="two"/><itemref idref="one"/></spine></package>'
    )
    with ZipFile(source, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OPS/content.opf", package)
        archive.writestr("OPS/first.xhtml", "<html><body><p>First chapter</p></body></html>")
        archive.writestr("OPS/second.xhtml", "<html><body><p>第二章</p></body></html>")


def test_epub_extracts_local_chapters_in_spine_order(tmp_path: Path) -> None:
    source = tmp_path / "book.epub"
    create_epub(source)
    output = tmp_path / "book.txt"
    convert_generated(source, "epub_text", output)
    text = output.read_text(encoding="utf-8")
    assert text.index("第二章") < text.index("First chapter")


@pytest.mark.parametrize("href", ["../../outside.xhtml", "https://invalid.example/chapter.xhtml"])
def test_epub_does_not_follow_escaped_or_external_chapter_paths(tmp_path: Path, href: str) -> None:
    source = tmp_path / "book.epub"
    create_epub(source, href)
    output = tmp_path / "book.txt"
    with pytest.raises(ConversionError, match="路径"):
        convert_generated(source, "epub_text", output)
    assert not output.exists()
    assert_no_staging(tmp_path)


def test_malformed_epub_reports_readable_error(tmp_path: Path) -> None:
    source = tmp_path / "broken.epub"
    with ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
    output = tmp_path / "broken.txt"
    with pytest.raises(ConversionError):
        convert_generated(source, "epub_text", output)
    assert not output.exists()
    assert_no_staging(tmp_path)


def create_wave(source: Path) -> None:
    """Write a brief silent PCM wave using only the standard library."""
    with wave.open(str(source), "wb") as stream:
        stream.setnchannels(1)
        stream.setsampwidth(2)
        stream.setframerate(8000)
        stream.writeframes(b"\0\0" * 2000)


def assert_media_decodes(binary: str, source: Path) -> str:
    """Decode generated output to a null sink, rejecting renamed containers."""
    result = subprocess.run(
        [binary, "-hide_banner", "-nostdin", "-i", str(source), "-f", "null", "-"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=30,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    assert result.returncode == 0, result.stderr
    return result.stderr


def test_ffmpeg_really_converts_generated_audio_and_video(tmp_path: Path) -> None:
    binary = media.ffmpeg_binary()
    assert binary, "Declared imageio-ffmpeg dependency must provide a local binary"
    audio = tmp_path / "tone.wav"
    create_wave(audio)
    audio_targets = {target.key for target in inspect_input(audio).targets}
    assert {"media_mp3", "media_flac"} <= audio_targets
    mp3 = tmp_path / "tone.mp3"
    convert_generated(audio, "media_mp3", mp3)
    assert "Audio: mp3" in assert_media_decodes(binary, mp3)
    flac = tmp_path / "tone.flac"
    convert_generated(audio, "media_flac", flac)
    assert flac.read_bytes().startswith(b"fLaC")
    assert "Audio: flac" in assert_media_decodes(binary, flac)
    video = tmp_path / "sample.mp4"
    generated = subprocess.run(
        [binary, "-hide_banner", "-nostdin", "-y", "-f", "lavfi", "-i", "color=c=blue:s=32x24:r=5:d=0.4",
         "-an", "-c:v", "libx264", "-pix_fmt", "yuv420p", str(video)],
        capture_output=True, timeout=30, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    assert generated.returncode == 0, generated.stderr.decode(errors="replace")
    assert "media_webm" in {target.key for target in inspect_input(video).targets}
    webm = tmp_path / "sample.webm"
    convert_generated(video, "media_webm", webm)
    assert "Video: vp9" in assert_media_decodes(binary, webm)
    assert webm.read_bytes().startswith(b"\x1a\x45\xdf\xa3")


def test_missing_backends_do_not_offer_unavailable_targets(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from filemind.conversion import office

    source = tmp_path / "sample.wav"
    create_wave(source)
    monkeypatch.setattr(media, "ffmpeg_binary", lambda: None)
    item = inspect_input(source)
    assert [target.key for target in item.targets] == ["zip"]
    assert "未安装" in item.description
    monkeypatch.setattr(office, "office_binary", lambda: None)
    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"generated placeholder")
    assert [target.key for target in inspect_input(legacy).targets] == ["zip"]
    with pytest.raises(ConversionError, match="不支持"):
        convert_generated(legacy, "lo_pdf", tmp_path / "legacy.pdf")


def test_invalid_media_is_rejected_without_a_fake_success(tmp_path: Path) -> None:
    source = tmp_path / "invalid.mp4"
    source.write_bytes(b"generated file without media streams")
    with pytest.raises(ConversionError, match="音视频流"):
        inspect_input(source)
    assert_no_staging(tmp_path)


def test_cancelling_ffmpeg_task_does_not_publish_or_leave_staging(tmp_path: Path) -> None:
    source = tmp_path / "audio.wav"
    create_wave(source)
    cancel = Event()

    def cancel_when_encoder_starts(_completed: int, total: int) -> None:
        if total == 0:
            cancel.set()

    output = tmp_path / "cancel.mp3"
    with pytest.raises(InterruptedError):
        convert_file(inspect_input(source), "media_mp3", output, cancel_when_encoder_starts, cancel)
    assert not output.exists()
    assert source.exists()
    assert_no_staging(tmp_path)


def test_extended_conversion_cancellation_and_output_conflict_preserve_files(tmp_path: Path) -> None:
    source = tmp_path / "rules.yaml"
    source.write_text("课程: 数据结构\n", encoding="utf-8")
    original = source.read_bytes()
    output = tmp_path / "rules.json"
    output.write_text("existing result", encoding="utf-8")
    with pytest.raises(ConversionError, match="不会覆盖"):
        convert_generated(source, "yaml_json", output)
    assert output.read_text(encoding="utf-8") == "existing result"
    cancel = Event()
    cancel.set()
    with pytest.raises(InterruptedError):
        convert_file(inspect_input(source), "yaml_json", tmp_path / "cancel.json", lambda *_: None, cancel)
    assert source.read_bytes() == original
    assert not (tmp_path / "cancel.json").exists()
    assert_no_staging(tmp_path)
