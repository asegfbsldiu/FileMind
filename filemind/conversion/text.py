"""Bounded text transformations; no scripts, external resources or commands."""

import html
import json
import posixpath
from html.parser import HTMLParser
from pathlib import Path
from threading import Event
from zipfile import ZipFile

from defusedxml import ElementTree

from filemind.conversion.base import ConversionError, Progress, check_archive, check_cancel
from filemind.conversion.models import ConversionTarget


TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".log", ".py", ".c", ".h", ".java", ".js", ".ts",
                 ".sql", ".csv", ".json", ".html", ".htm", ".xml", ".yaml", ".yml"}
EXTRA_TARGETS = {
    "text_epub": ConversionTarget("text_epub", "文本 → EPUB 电子书", ".epub", "创建含目录的 EPUB 3，按文本段落排版；不解析脚本或复刻原排版。"),
    "text_html": ConversionTarget("text_html", "文本 → 静态 HTML", ".html", "转义文本，保留换行；不执行脚本或加载远程资源。"),
    "html_text": ConversionTarget("html_text", "HTML → 纯文本", ".txt", "提取可见文本，忽略脚本/样式；不加载外部链接。"),
    "text_docx": ConversionTarget("text_docx", "文本 → DOCX", ".docx", "按文本行创建段落，不承诺 Markdown 或 HTML 排版复刻。"),
    "json_yaml": ConversionTarget("json_yaml", "JSON → YAML", ".yaml", "使用本机安全序列化，不执行 YAML 对象构造器。"),
    "yaml_json": ConversionTarget("yaml_json", "YAML → JSON", ".json", "安全加载单文档 YAML；仅支持 JSON 可表达的数据，不执行标签。"),
    "xml_pretty": ConversionTarget("xml_pretty", "格式化 XML", ".xml", "禁用 DTD 和实体；格式化可能改变前缀和空白，不适用于数字签名 XML。"),
    "csv_xlsx": ConversionTarget("csv_xlsx", "CSV → XLSX", ".xlsx", "所有单元格按文本写入，不执行以等号开头的公式。"),
}


def extra_text_targets(suffix: str) -> tuple[ConversionTarget, ...]:
    """Expose transformations matching this text format, keeping original IDs."""
    keys = ["text_html", "text_docx", "text_epub"]
    if suffix in {".html", ".htm"}:
        keys = ["html_text", "text_docx", "text_epub"]
    keys += {".json": ["json_yaml"], ".yaml": ["yaml_json"], ".yml": ["yaml_json"],
             ".xml": ["xml_pretty"], ".csv": ["csv_xlsx"]}.get(suffix, [])
    return tuple(EXTRA_TARGETS[key] for key in keys)


class VisibleText(HTMLParser):
    """Collect text without fetching images, scripts, CSS or linked documents."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.hidden = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag in {"p", "div", "br", "h1", "h2", "li", "tr"} and not self.hidden:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.hidden:
            self.hidden -= 1

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.parts.append(data)


def html_text(text: str) -> str:
    """Strip active content while retaining readable text boundaries."""
    parser = VisibleText()
    parser.feed(text)
    return "".join(parser.parts).strip()


def write_transformed(text: str, key: str, staging: Path, cancel: Event) -> None:
    """Write one real transformation to a caller-owned staging file."""
    check_cancel(cancel)
    if key == "text_epub":
        from filemind.conversion.epub import write_epub
        write_epub(html_text(text) if text.lstrip().lower().startswith(("<!doctype html", "<html")) else text, staging, cancel)
        return
    if key == "text_docx":
        from docx import Document
        document = Document()
        for line in text.splitlines():
            check_cancel(cancel)
            document.add_paragraph(line)
        document.save(staging)
        return
    if key == "csv_xlsx":
        _csv_workbook(text, staging, cancel)
        return
    if key == "text_html":
        result = f'<!doctype html><html><meta charset="utf-8"><body><pre>{html.escape(text)}</pre></body></html>'
    elif key == "html_text":
        result = html_text(text)
    elif key == "json_yaml":
        import yaml
        result = yaml.safe_dump(json.loads(text), allow_unicode=True, sort_keys=False)
    elif key == "yaml_json":
        result = _yaml_json(text)
    elif key == "xml_pretty":
        result = _format_xml(text)
    else:
        raise ConversionError("不支持这个文本转换目标。")
    staging.write_text(result, encoding="utf-8")


def _format_xml(text: str) -> str:
    from defusedxml.common import DefusedXmlException
    from xml.etree.ElementTree import ParseError, indent, tostring
    try:
        root = ElementTree.fromstring(text, forbid_dtd=True, forbid_entities=True, forbid_external=True)
        indent(root)
        return tostring(root, encoding="unicode")
    except (DefusedXmlException, ParseError) as error:
        raise ConversionError("XML 格式无效或包含被禁用的 DTD / 实体，未读取外部资源。") from error


def _yaml_json(text: str) -> str:
    import yaml
    try:
        if any(isinstance(event, yaml.AliasEvent) for event in yaml.parse(text)):
            raise ConversionError("暂不转换含 YAML 别名的数据，避免循环或过度展开。")
        value = yaml.safe_load(text)
    except yaml.YAMLError as error:
        raise ConversionError("YAML 格式无效或包含不支持的标签，未执行任何内容。") from error
    try:
        return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)
    except (ValueError, TypeError) as error:
        raise ConversionError("YAML 包含日期、非有限数等 JSON 无法直接表达的数据。") from error


def _csv_workbook(text: str, output: Path, cancel: Event) -> None:
    import csv
    import io
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    for index, row in enumerate(csv.reader(io.StringIO(text))):
        check_cancel(cancel)
        if index >= 50000 or len(row) > 200:
            raise ConversionError("CSV 超过 50000 行 / 200 列限制。")
        for column, value in enumerate(row, 1):
            cell = sheet.cell(index + 1, column, value)
            cell.data_type = "s"
    workbook.save(output)
    workbook.close()


def convert_epub(source: Path, staging: Path, progress: Progress, cancel: Event) -> None:
    """Read only bounded spine entries from the selected EPUB archive."""
    from xml.etree.ElementTree import ParseError
    from defusedxml.common import DefusedXmlException
    try:
        _write_epub(source, staging, progress, cancel)
    except (KeyError, AttributeError, ParseError, DefusedXmlException, UnicodeError) as error:
        raise ConversionError("EPUB 章节索引或文本无效，可能是加密、损坏或非 UTF-8 内容。") from error


def _write_epub(source: Path, staging: Path, progress: Progress, cancel: Event) -> None:
    check_archive(source)
    with ZipFile(source) as archive:
        container = ElementTree.fromstring(archive.read("META-INF/container.xml"))
        rootfile = container.find(".//{*}rootfile")
        package_path = rootfile.attrib["full-path"]
        package = ElementTree.fromstring(archive.read(package_path))
        manifest = {item.attrib["id"]: item.attrib["href"] for item in package.findall(".//{*}manifest/{*}item")}
        spine = package.findall(".//{*}spine/{*}itemref")
        with staging.open("w", encoding="utf-8") as output:
            for index, item in enumerate(spine):
                check_cancel(cancel)
                relative = manifest[item.attrib["idref"]].split("#", 1)[0]
                entry = posixpath.normpath(posixpath.join(posixpath.dirname(package_path), relative))
                if entry.startswith(("/", "../")) or ":" in entry:
                    raise ConversionError("EPUB 包含不合法的章节路径。")
                output.write(html_text(archive.read(entry).decode("utf-8-sig")) + "\n\n")
                progress(index + 1, len(spine))
