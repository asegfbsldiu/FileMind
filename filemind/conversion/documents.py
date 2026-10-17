"""Office text/data adapters; document rendering is a separate optional backend."""

import csv
import html
import io
import json
from pathlib import Path
from threading import Event
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile

from filemind.conversion.base import ConversionError, Progress, check_archive, check_cancel
from filemind.conversion.models import ConversionTarget


DOCUMENT_TEXT = ConversionTarget("office_text", "提取文字 → TXT", ".txt", "仅提取段落/表格文字，不保留页面布局、图片和公式渲染。")
DOCUMENT_HTML = ConversionTarget("office_html", "文字内容 → HTML", ".html", "离线静态文字页面；不保留原文档排版，不包含外链或脚本。")
SHEET_JSON = ConversionTarget("sheet_json", "全部工作表 → JSON", ".json", "导出值；公式只使用文件内已有缓存值，不计算、不连接外部数据。")
SHEET_CSV = ConversionTarget("sheet_csv", "全部工作表 → CSV ZIP", ".zip", "每张工作表一个 UTF-8 CSV；不保留格式，不执行公式。")


def inspect_document(source: Path) -> tuple[str, str, tuple[ConversionTarget, ...]] | None:
    """Check package structure before offering a content conversion."""
    if source.suffix.lower() not in {".docx", ".xlsx", ".pptx"}:
        return None
    try:
        check_archive(source)
        with ZipFile(source) as archive:
            names = set(archive.namelist())
        if "word/document.xml" in names and source.suffix.lower() == ".docx":
            return "document", "DOCX · 文字提取（非排版转换）", (DOCUMENT_TEXT, DOCUMENT_HTML)
        if "xl/workbook.xml" in names and source.suffix.lower() == ".xlsx":
            return "sheet", "XLSX · 全工作表值导出", (SHEET_JSON, SHEET_CSV)
        if "ppt/presentation.xml" in names and source.suffix.lower() == ".pptx":
            return "slides", "PPTX · 幻灯片文字提取", (DOCUMENT_TEXT, DOCUMENT_HTML)
    except BadZipFile:
        return None
    return None


def _document_lines(source: Path, cancel: Event) -> list[str]:
    if source.suffix.lower() == ".docx":
        from docx import Document
        document = Document(source)
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            check_cancel(cancel)
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return lines
    from pptx import Presentation
    presentation = Presentation(source)
    lines = []
    for index, slide in enumerate(presentation.slides):
        check_cancel(cancel)
        lines.append(f"第 {index + 1} 页")
        lines.extend(shape.text for shape in slide.shapes if shape.has_text_frame)
    return lines


def convert_document(source: Path, key: str, staging: Path, progress: Progress, cancel: Event) -> None:
    """Extract visible text or workbook values without invoking Office macros."""
    check_archive(source)
    if key.startswith("sheet_"):
        _convert_workbook(source, key, staging, progress, cancel)
        return
    lines = _document_lines(source, cancel)
    check_cancel(cancel)
    if key == "office_html":
        body = "\n".join(f"<p>{html.escape(line)}</p>" for line in lines)
        text = f'<!doctype html><html><head><meta charset="utf-8"></head><body>{body}</body></html>'
    else:
        text = "\n".join(lines)
    staging.write_text(text, encoding="utf-8")
    progress(1, 1)


def _convert_workbook(source: Path, key: str, staging: Path, progress: Progress, cancel: Event) -> None:
    from openpyxl import load_workbook
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        sheets: dict[str, list[list[object]]] = {}
        cells = 0
        for index, sheet in enumerate(workbook):
            if (sheet.max_row or 0) > 50000 or (sheet.max_column or 0) > 200:
                raise ConversionError("表格声明的尺寸超过每表 50000 行 / 200 列限制。")
            rows = []
            for row in sheet.iter_rows(values_only=True):
                check_cancel(cancel)
                cells += len(row)
                if len(rows) >= 50000 or len(row) > 200 or cells > 500000:
                    raise ConversionError("表格超过行列或总计 50 万单元格限制。")
                rows.append(list(row))
            sheets[sheet.title] = rows
            progress(index + 1, len(workbook.worksheets))
        if key == "sheet_json":
            staging.write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        else:
            _write_sheet_archive(sheets, staging)
    finally:
        workbook.close()


def _write_sheet_archive(sheets: dict[str, list[list[object]]], staging: Path) -> None:
    with ZipFile(staging, "w", compression=ZIP_DEFLATED) as archive:
        for index, (name, rows) in enumerate(sheets.items(), 1):
            with io.StringIO(newline="") as stream:
                csv.writer(stream).writerows(rows)
                safe_name = "".join(character if character.isalnum() else "_" for character in name)[:60]
                archive.writestr(f"{index:02}-{safe_name}.csv", stream.getvalue().encode("utf-8-sig"))
