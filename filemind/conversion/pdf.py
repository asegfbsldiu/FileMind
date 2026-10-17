"""PDFium text extraction and page-image export, serialized across threads."""

import io
import pypdfium2 as pdfium
from pathlib import Path
from threading import Event, RLock
from zipfile import ZIP_DEFLATED, ZipFile

from filemind.conversion.base import ConversionError, Progress, check_cancel
from filemind.conversion.models import ConversionTarget


PDF_LOCK = RLock()
TARGETS = (
    ConversionTarget("pdf_docx", "PDF 文本 → DOCX", ".docx", "仅提取文本层重建段落，不保留原排版、图片或表格结构；不执行 OCR。"),
    ConversionTarget("pdf_html", "PDF 文本 → 静态 HTML", ".html", "按页导出转义的文本层，不保留原排版和图片；不执行 OCR。"),
    ConversionTarget("pdf_text", "PDF → TXT（文本层）", ".txt", "按页提取文本层，不执行 OCR；扫描版 PDF 可能无文本。"),
    ConversionTarget("pdf_png", "PDF → PNG 页面 ZIP", ".zip", "逐页渲染为 144 DPI PNG，再打包；最多 100 页，不承诺保留交互表单。"),
    ConversionTarget("pdf_jpeg", "PDF → JPEG 页面 ZIP", ".zip", "逐页渲染并打包 JPEG；有损压缩，最多 100 页。"),
)


def inspect_pdf(source: Path) -> tuple[str, str, tuple[ConversionTarget, ...]]:
    """Inspect page count under the same mutex used by all PDFium calls."""
    try:
        with PDF_LOCK, pdfium.PdfDocument(source) as document:
            count = len(document)
            if count < 1 or count > 100:
                raise ConversionError("PDF 页数必须为 1–100 页。")
            return "pdf", f"PDF · {count} 页 · 本地 PDFium", TARGETS
    except pdfium.PdfiumError as error:
        raise ConversionError("PDF 无法打开，可能已加密或损坏；请先在本机解密。") from error


def convert_pdf(source: Path, key: str, staging: Path, progress: Progress, cancel: Event) -> None:
    """Release each PDF page before continuing, keeping memory bounded."""
    with PDF_LOCK, pdfium.PdfDocument(source) as document:
        if key in {"pdf_docx", "pdf_html"}:
            _export_text_document(document, key, staging, progress, cancel)
        elif key == "pdf_text":
            _extract_text(document, staging, progress, cancel)
        else:
            _export_pages(document, key, staging, progress, cancel)


def _export_text_document(document: pdfium.PdfDocument, key: str, output: Path, progress: Progress, cancel: Event) -> None:
    from html import escape
    from docx import Document

    pages: list[str] = []
    total = 0
    for index in range(len(document)):
        check_cancel(cancel)
        page = document[index]
        try:
            textpage = page.get_textpage()
            try:
                if textpage.count_chars() + total > 2_000_000:
                    raise ConversionError("PDF 文本超过 200 万字符限制。")
                text = textpage.get_text_bounded()
                total += len(text)
                pages.append(text)
            finally:
                textpage.close()
        finally:
            page.close()
        progress(index + 1, len(document))
    if not any(text.strip() for text in pages):
        raise ConversionError("PDF 没有可提取的文本层；未执行 OCR。")
    if key == "pdf_html":
        body = "".join(f"<section><h2>第 {index + 1} 页</h2><pre>{escape(text)}</pre></section>" for index, text in enumerate(pages))
        output.write_text(f'<!doctype html><html><meta charset="utf-8"><body>{body}</body></html>', encoding="utf-8")
    else:
        result = Document()
        for index, text in enumerate(pages):
            check_cancel(cancel)
            if index:
                result.add_page_break()
            for line in text.splitlines():
                result.add_paragraph(line)
        result.save(output)


def _extract_text(document: pdfium.PdfDocument, output: Path, progress: Progress, cancel: Event) -> None:
    nonempty = False
    with output.open("w", encoding="utf-8") as stream:
        for index in range(len(document)):
            check_cancel(cancel)
            page = document[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_bounded()
                finally:
                    textpage.close()
                nonempty = nonempty or bool(text.strip())
                stream.write(f"--- 第 {index + 1} 页 ---\n{text}\n\n")
            finally:
                page.close()
            progress(index + 1, len(document))
    if not nonempty:
        raise ConversionError("PDF 没有可提取的文本层；未执行 OCR，可改用页面图像导出。")


def _export_pages(document: pdfium.PdfDocument, key: str, output: Path, progress: Progress, cancel: Event) -> None:
    image_format, extension = ("PNG", "png") if key == "pdf_png" else ("JPEG", "jpg")
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for index in range(len(document)):
            check_cancel(cancel)
            page = document[index]
            try:
                width, height = page.get_size()
                if width * height * 4 > 20_000_000:
                    raise ConversionError("PDF 单页渲染超过 2000 万像素限制。")
                bitmap = page.render(scale=2)
                try:
                    with bitmap.to_pil() as image, io.BytesIO() as buffer:
                        image.convert("RGB").save(buffer, format=image_format)
                        archive.writestr(f"page-{index + 1:03}.{extension}", buffer.getvalue())
                finally:
                    bitmap.close()
            finally:
                page.close()
            progress(index + 1, len(document))
