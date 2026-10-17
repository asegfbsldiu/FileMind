"""Offline release diagnostics using generated samples, never user application state."""

from __future__ import annotations

import argparse
import importlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Callable

from filemind import __version__


def _check_ui() -> dict[str, object]:
    from PySide6.QtWidgets import QApplication
    from filemind.ui.branding import application_icon
    from filemind.ui.styles.theme import apply_theme

    application = QApplication.instance() or QApplication([])
    apply_theme(application)
    for module in ("main_window", "pages.organize_page", "pages.assistant_page",
                   "pages.conversion_page", "pages.settings_page", "pages.local_settings_page"):
        importlib.import_module(f"filemind.ui.{module}")
    icon = application_icon()
    if icon.isNull() or icon.pixmap(32, 32).isNull():
        raise RuntimeError("Application icon is unavailable")
    stylesheet = application.styleSheet()
    if "__card__" in stylesheet or "#OrbitRoot" not in stylesheet:
        raise RuntimeError("Bundled stylesheets are unavailable")
    return {"platform": application.platformName(), "icon": True, "qss": True}


def _check_documents(directory: Path) -> dict[str, object]:
    from docx import Document
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation

    document = Document()
    document.add_paragraph("FileMind generated release sample")
    word_path = directory / "sample.docx"
    document.save(word_path)
    if not Document(word_path).paragraphs:
        raise RuntimeError("DOCX template/readback failed")
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    ppt_path = directory / "sample.pptx"
    presentation.save(ppt_path)
    if len(Presentation(ppt_path).slides) != 1:
        raise RuntimeError("PPTX template/readback failed")
    book = Workbook()
    book.active["A1"] = "FileMind"
    sheet_path = directory / "sample.xlsx"
    book.save(sheet_path)
    book.close()
    workbook = load_workbook(sheet_path)
    try:
        if workbook.active["A1"].value != "FileMind":
            raise RuntimeError("XLSX readback failed")
    finally:
        workbook.close()
    return {"docx": True, "pptx": True, "xlsx": True}


def _check_images_pdf(directory: Path) -> dict[str, object]:
    from PIL import Image
    import pypdfium2 as pdfium

    image_path = directory / "sample.png"
    with Image.new("RGB", (8, 8), "#7458E8") as image:
        image.save(image_path)
    with Image.open(image_path) as image:
        pdf_path = directory / "sample.pdf"
        image.save(pdf_path, "PDF")
    with pdfium.PdfDocument(pdf_path) as document:
        if len(document) != 1:
            raise RuntimeError("PDFium readback failed")
        page = document[0]
        bitmap = page.render(scale=1)
        try:
            rendered = bitmap.to_pil()
            rendered.close()
        finally:
            bitmap.close()
            page.close()
    return {"pillow_png": True, "pillow_pdf": True, "pdfium_render": True}


def _check_fonts() -> dict[str, object]:
    import brotli
    from fontTools.ttLib import woff2

    sample = b"FileMind generated font dependency check"
    if brotli.decompress(brotli.compress(sample)) != sample or not callable(woff2.compress):
        raise RuntimeError("Font/Brotli dependencies are unavailable")
    return {"fonttools_woff2": True, "brotli": True}


def _check_ffmpeg() -> dict[str, object]:
    import imageio_ffmpeg

    directory = Path(imageio_ffmpeg.__file__).parent / "binaries"
    candidates = sorted(directory.glob("ffmpeg*.exe"))
    if len(candidates) != 1:
        raise RuntimeError("A unique bundled FFmpeg executable was not found")
    command = [str(candidates[0]), "-hide_banner", "-L"]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8",
                            errors="replace", timeout=12, check=True,
                            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    license_text = (result.stdout + result.stderr).strip()
    if "license" not in license_text.lower():
        raise RuntimeError("Bundled FFmpeg license output is unavailable")
    return {"bundled_ffmpeg": True, "ffmpeg_license": license_text}


def run_checks() -> dict[str, object]:
    """Exercise only generated inputs and report each failing dependency group."""
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    result: dict[str, object] = {"version": __version__, "frozen": bool(getattr(sys, "frozen", False)),
                                "ok": True, "checks": {}}
    checks: dict[str, object] = {}
    with tempfile.TemporaryDirectory(prefix="filemind-release-check-") as temporary:
        directory = Path(temporary)
        operations: dict[str, Callable[[], dict[str, object]]] = {
            "ui": _check_ui, "documents": lambda: _check_documents(directory),
            "images_pdf": lambda: _check_images_pdf(directory), "fonts": _check_fonts,
            "ffmpeg": _check_ffmpeg,
        }
        for name, operation in operations.items():
            try:
                details = operation()
                if "ffmpeg_license" in details:
                    result["ffmpeg_license"] = details.pop("ffmpeg_license")
                checks[name] = {"ok": True, **details}
            except (OSError, RuntimeError, ValueError, ImportError, subprocess.SubprocessError) as error:
                checks[name] = {"ok": False, "error_type": type(error).__name__, "error": str(error)}
                result["ok"] = False
    result["checks"] = checks
    return result


def main(arguments: list[str] | None = None) -> int:
    """Write diagnostics to an explicit report file, including in GUI builds."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--self-check", action="store_true", required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args(arguments)
    result = run_checks()
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", "utf-8")
    return 0 if result["ok"] else 1
